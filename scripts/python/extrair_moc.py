#!/usr/bin/env python3
"""
extrair_moc.py - Extrai informacoes estruturadas de MOCs (PDF e DOCX).

Uso:
    python extrair_moc.py <arquivo.pdf|docx> [--manifest <path>]

Dependencias:
    pip install pymupdf>=1.24.0 pyyaml python-docx>=1.1.0
"""

import sys
import os
import re
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pymupdf

from utils.cores import extract_page_content, aggregate_spans_into_lines
from utils.parsing import (
    parse_moc_metadata,
    calculate_sha256,
)
from utils.moc import montar_moc_estrutura
from utils.saida import gerar_json_moc, salvar_json
from utils.manifest import adicionar_manifest_moc


def itens_from_pdf(pdf_path):
    """Extrai linhas do PDF com marcacao/pagina."""
    doc = pymupdf.open(pdf_path)
    all_items = []
    page_texts = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        raw_items = extract_page_content(page)
        page_items = aggregate_spans_into_lines(raw_items)
        text = page.get_text()
        page_texts.append(text)
        for item in page_items:
            item["pagina"] = page_num + 1
        all_items.extend(page_items)
    num_paginas = len(doc)
    doc.close()
    return all_items, page_texts, num_paginas


def itens_from_docx(docx_path):
    """Extrai linhas do DOCX (sem cor: SEM_MARCA), incluindo celulas de tabelas."""
    try:
        import docx as python_docx
    except ImportError:
        print("ERRO: python-docx nao instalado. Execute: pip install python-docx", file=sys.stderr)
        return [], [], 0

    doc = python_docx.Document(docx_path)
    all_items = []
    page_texts = []
    current_page = []
    pagina = 1

    def _emitir(texto):
        nonlocal pagina
        if texto is None:
            return
        texto = texto.replace("\t", " ")
        for linha in texto.splitlines():
            s = linha.strip()
            if s:
                all_items.append({"text": s, "marcacao": "SEM_MARCA", "pagina": pagina})
            current_page.append(linha)
            if len(current_page) >= 50:
                page_texts.append("\n".join(current_page))
                current_page.clear()
                pagina += 1

    def _walk_tables(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        _emitir(p.text)
                    _walk_tables(cell.tables)

    for para in doc.paragraphs:
        _emitir(para.text)
    _walk_tables(doc.tables)

    if current_page:
        page_texts.append("\n".join(current_page))

    return all_items, page_texts, pagina


def filtrar_ruido(all_items, num_paginas):
    """Remove linhas curtas, paginacao, cabecalhos repetitivos e pontilhados."""
    itens = []
    for item in all_items:
        text = item.get("text", "").strip()
        if len(text) < 3:
            continue
        if re.match(r"^Página \d+ / \d+$", text):
            continue
        if re.match(r"^[\.\-•—]{16,}$", text):
            continue
        if text in ("•", "", ">", ">>", "-"):
            continue
        if re.search(r"\.{15,}\s*\d*$", text):
            continue
        itens.append(item)

    # Cabecalhos/rodapes que se repetem em varias paginas
    paginas_por_texto = {}
    for item in itens:
        t = item.get("text", "").strip()
        paginas_por_texto.setdefault(t, set()).add(item.get("pagina"))
    min_repeticoes = max(3, int(num_paginas * 0.4)) if num_paginas >= 3 else 3
    itens = [
        item
        for item in itens
        if item.get("text", "").strip().startswith("#")
        or len(paginas_por_texto.get(item.get("text", "").strip(), ())) < min_repeticoes
    ]
    return itens


def extrair_moc(file_path, manifest_path=None):
    """
    Orquestra extracao completa de um MOC.

    Args:
        file_path: caminho para o arquivo (PDF ou DOCX)
        manifest_path: caminho do manifest.yaml (opcional)

    Returns:
        dict com resultado da extracao
    """
    file_path = os.path.abspath(file_path)
    filename = os.path.basename(file_path)
    ext = Path(filename).suffix.lower()

    print(f"Processando MOC: {filename}")

    if ext == ".docx":
        all_items, page_texts, num_paginas = itens_from_docx(file_path)
    else:
        all_items, page_texts, num_paginas = itens_from_pdf(file_path)

    if not all_items:
        return {"erro": "Nao foi possivel extrair conteudo do arquivo"}

    metadata = parse_moc_metadata(page_texts, filename)
    print(f"Documento: {metadata['documento']}, Versao: {metadata['versao']}")
    print(f"Secao: {metadata['secao']}")

    linhas = filtrar_ruido(all_items, num_paginas)

    # MOCs em PDF podem numerar secoes como "N. Titulo" (ex: Anexo II DACTE
    # do CT-e, "1. Introducao"). Habilita o formato "com ponto" quando ha
    # pelo menos 3 linhas com essa forma; DOCX sempre usa o formato.
    permitir_numero_ponto = ext == ".docx"
    if not permitir_numero_ponto:
        contagem = 0
        for item in linhas:
            t = item.get("text", "").strip()
            if re.match(r"^\d{1,2}\.\s+[A-ZÀ-Ú]", t) or re.match(
                r"^\d+(?:\.\d+)+\.\s+[A-ZÀ-Ú]", t
            ):
                contagem += 1
                if contagem >= 3:
                    permitir_numero_ponto = True
                    break

    resultado = montar_moc_estrutura(linhas, permitir_numero_ponto=permitir_numero_ponto)
    secoes = resultado["secoes"]
    regras = resultado["regras_validacao"]
    campos = resultado["campos_leiaute"]
    stats = resultado["estatisticas"]

    print(f"Secoes: {len(secoes)} | Regras: {len(regras)} | Campos: {len(campos)}")

    sha256 = calculate_sha256(file_path)

    json_data = gerar_json_moc(
        metadata=metadata,
        secoes=secoes,
        regras_validacao=regras,
        campos_leiaute=campos,
        arquivo_origem=file_path,
        sha256=sha256,
    )
    json_data["estatisticas"] = stats

    doc_lower = metadata["documento"].lower().replace("-", "").replace("e", "")
    if doc_lower == "nf":
        doc_lower = "nfe"
    elif doc_lower == "ct":
        doc_lower = "cte"
    elif doc_lower == "mdf":
        doc_lower = "mdfe"
    output_dir = os.path.join("catalogo", "moc", doc_lower)

    json_filename = f"{Path(filename).stem}.json"
    json_path = salvar_json(json_data, output_dir, json_filename)
    print(f"JSON: {json_path}")

    if manifest_path is None:
        manifest_path = "manifest.yaml"

    if os.path.exists(manifest_path):
        adicionar_manifest_moc(
            {
                "documento": metadata["documento"],
                "versao": metadata["versao"],
                "secao": metadata.get("secao", ""),
                "arquivo_origem": os.path.relpath(file_path, os.path.dirname(manifest_path) or "."),
                "sha256": sha256,
            },
            manifest_path,
        )

    return json_data


def main():
    """CLI entry point."""
    import argparse

    parser = argparse.ArgumentParser(
        description="Extrai informacoes de MOCs (PDF ou DOCX)"
    )
    parser.add_argument("arquivo", help="Caminho para o arquivo do MOC")
    parser.add_argument(
        "--manifest",
        default="manifest.yaml",
        help="Caminho para manifest.yaml",
    )

    args = parser.parse_args()

    if not os.path.exists(args.arquivo):
        print(f"ERRO: Arquivo nao encontrado: {args.arquivo}", file=sys.stderr)
        sys.exit(1)

    result = extrair_moc(args.arquivo, args.manifest)

    if "erro" in result:
        print(f"ERRO: {result['erro']}", file=sys.stderr)
        sys.exit(1)

    print(f"\nConcluido: MOC {result['documento']} v{result['versao']}")
    print(f"Secoes: {len(result.get('secoes', []))}")
    print(f"Regras: {len(result.get('regras_validacao', []))}")
    print(f"Campos: {len(result.get('campos_leiaute', []))}")


if __name__ == "__main__":
    main()