#!/usr/bin/env python3
"""
verificar_extracao.py - Confere JSONs extraidos (NT/MOC) contra o PDF/DOCX de origem.

Uso:
    python verificar_extracao.py                      # todos os JSONs de catalogo/nt e catalogo/moc
    python verificar_extracao.py <arquivo.json> ...   # arquivos especificos
    python verificar_extracao.py --so-bloqueantes     # esconde avisos

Niveis:
    BLOQUEANTE - extracao nao confiavel; codigo de saida 1
    AVISO      - suspeito, revisar

Somente le `fontes/` (regra 11 do AGENTS.md).
"""

import collections
import glob
import json
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pymupdf

from utils.json_schema import validate_nt_json, validate_moc_json

RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

BLOQUEANTE = "BLOQUEANTE"
AVISO = "AVISO"

# ID de regra de NF-e impresso sozinho no inicio da linha (ex: B25-80, 5AF15-11, I11c-10, 12C02-10)
PAT_ID_LINHA = re.compile(r"(?m)^\s*([0-9]{0,2}[A-Z]{1,3}[0-9]{2}[a-z]?-[0-9]{2,3})\s*$")
# Mensagem que termina em palavra de ligacao = quase certamente cortada
# (so minusculas: a sigla 'FS-DA' termina em 'DA' e nao e preposicao)
# e nao e o 'e' de sigla com hifen ('NF-e', 'NFC-e', 'DF-e')
PAT_TERMINA_LIGACAO = re.compile(
    r"(?<![\w-])(do|da|de|dos|das|na|no|nas|nos|para|com|e|ou|em|ao|a|o|que|por|sem|se|pelo|pela)\s*$"
)
# Nota "Observacao/Excecao" com fragmento colado depois do ponto final: "... Autorizadoras. irregular na UF"
# (assinatura do defeito de mistura de colunas; linhas de condicao como "...; e" NAO sao fragmento)
PAT_NOTA_COLADA = re.compile(r"^(Observa[cç][aã]o|Exce[cç][aã]o)\b.*[.]\s+[a-zà-ú]")
PAT_VERSAO_VALIDA = re.compile(r"^\d+(?:\.\d+)*[a-z]?$")
PAT_DATA_INVALIDA = re.compile(r"00/00/0000")


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------
def _palavras(texto):
    return re.findall(r"\w+", texto.lower())


def _coletar_texto(obj, acc):
    if isinstance(obj, str):
        acc.append(obj)
    elif isinstance(obj, dict):
        for v in obj.values():
            _coletar_texto(v, acc)
    elif isinstance(obj, list):
        for v in obj:
            _coletar_texto(v, acc)


def resolver_origem(arquivo_origem):
    """Resolve arquivo_origem (relativo a raiz ou absoluto) para um caminho existente, ou None."""
    if not arquivo_origem:
        return None
    candidatos = [arquivo_origem, os.path.join(RAIZ, arquivo_origem)]
    base = os.path.basename(arquivo_origem.replace("\\", "/"))
    candidatos += glob.glob(os.path.join(RAIZ, "fontes", "*", "*", base))
    for c in candidatos:
        if os.path.isfile(c):
            return c
    return None


def texto_do_documento(caminho):
    """Texto bruto (PDF via PyMuPDF; DOCX via python-docx) e quantidade de paginas/paragrafos."""
    if caminho.lower().endswith(".docx"):
        import docx

        d = docx.Document(caminho)
        partes = [p.text for p in d.paragraphs]
        for tb in d.tables:
            for r in tb.rows:
                for c in r.cells:
                    partes.append(c.text)
        return "\n".join(partes), len(d.paragraphs)
    doc = pymupdf.open(caminho)
    try:
        return "\n".join(p.get_text() for p in doc), len(doc)
    finally:
        doc.close()


def fundos_coloridos(caminho):
    """Conta paginas com retangulos de fundo amarelo/verde no PDF (indicio de marcacoes)."""
    if not caminho.lower().endswith(".pdf"):
        return {"AMARELO": set(), "VERDE": set()}
    cores = {(1.0, 1.0, 0.0): "AMARELO", (0.0, 1.0, 0.0): "VERDE"}
    achados = {"AMARELO": set(), "VERDE": set()}
    doc = pymupdf.open(caminho)
    try:
        for i, pg in enumerate(doc, 1):
            for d in pg.get_drawings():
                f = d.get("fill")
                if f:
                    nome = cores.get(tuple(round(x, 2) for x in f))
                    if nome:
                        achados[nome].add(i)
    finally:
        doc.close()
    return achados


def regras_da_nt(dados):
    """Lista de regras (dict) de todas as secoes 'regras' de uma NT."""
    saida = []
    for s in dados.get("secoes", []):
        if s.get("tipo") == "regras":
            saida.extend(s.get("regras", []))
    return saida


def mensagem_truncada(msg, colchetes=True):
    """
    True se a mensagem parece cortada: termina em palavra de ligacao (minuscula) ou, se `colchetes`,
    tem colchetes desbalanceados. Nos MOCs use colchetes=False: alguns exemplos do proprio PDF
    ('[Param: xxx)') sao desbalanceados.
    """
    msg = re.sub(r"-\s+", "-", (msg or "").strip())  # 'NF- e' (espaco do PDF) -> 'NF-e'
    if not msg:
        return False
    if PAT_TERMINA_LIGACAO.search(msg):
        return True
    return colchetes and msg.count("[") != msg.count("]")


def cronograma_invalido(cronograma):
    """Lista de motivos pelos quais um cronograma tem valores invalidos."""
    motivos = []
    for c in cronograma:
        v = str(c.get("versao", ""))
        if not PAT_VERSAO_VALIDA.match(v):
            motivos.append(f"versao invalida {v!r}")
        for campo in ("homologacao", "producao"):
            if PAT_DATA_INVALIDA.search(str(c.get(campo, ""))):
                motivos.append(f"{campo} com 00/00/0000")
    return list(dict.fromkeys(motivos))


# ---------------------------------------------------------------------------
# Verificacao
# ---------------------------------------------------------------------------
def verificar_json(caminho_json):
    """
    Confere um JSON extraido.

    Returns:
        dict com 'arquivo', 'tipo', 'problemas' [(nivel, codigo, texto)] e 'info'
    """
    with open(caminho_json, encoding="utf-8") as f:
        dados = json.load(f)
    nome = os.path.splitext(os.path.basename(caminho_json))[0]
    tipo = dados.get("tipo_documento", "NT")
    problemas = []
    info = {}

    def add(nivel, codigo, texto):
        problemas.append((nivel, codigo, texto))

    # 1. Schema
    validador = validate_nt_json if tipo == "NT" else validate_moc_json
    ok, erros = validador(dados)
    if not ok:
        add(BLOQUEANTE, "SCHEMA", "; ".join(map(str, erros))[:300])

    # 2. Origem portavel e existente
    origem = dados.get("arquivo_origem", "")
    if re.match(r"^[A-Za-z]:[\\/]", origem) or origem.startswith(("/", "\\")):
        add(BLOQUEANTE, "ORIGEM_ABSOLUTA", f"arquivo_origem absoluto: {origem}")
    caminho_origem = resolver_origem(origem)
    if caminho_origem is None:
        add(BLOQUEANTE, "ORIGEM_AUSENTE", f"documento de origem nao encontrado: {origem}")

    # 3. Metadados (NT)
    if tipo == "NT":
        titulo = dados.get("titulo", "")
        if not titulo or titulo.startswith("Projeto ") or len(titulo) < 12:
            add(BLOQUEANTE, "TITULO_SUSPEITO", f"titulo suspeito: {titulo!r}")
        m = re.search(r"_v(\d+(?:\.\d+)*[a-z]?)", os.path.basename(origem.replace("\\", "/")))
        if m and m.group(1) != dados.get("versao"):
            add(BLOQUEANTE, "VERSAO_DIVERGENTE",
                f"versao do JSON {dados.get('versao')!r} difere do nome do arquivo {m.group(1)!r}")

    # 4. Cobertura de texto
    if caminho_origem:
        texto, tam = texto_do_documento(caminho_origem)
        pw = _palavras(texto)
        acc = []
        _coletar_texto(dados, acc)
        jw = collections.Counter(_palavras(" ".join(acc)))
        faltando = [w for w in pw if jw[w] == 0]
        cobertura = 1 - len(faltando) / max(1, len(pw))
        info["cobertura_texto"] = round(cobertura, 4)
        if cobertura < 0.90:
            add(BLOQUEANTE, "COBERTURA_BAIXA", f"cobertura de texto {cobertura:.1%}")
        elif cobertura < 0.95:
            add(AVISO, "COBERTURA_MEDIA", f"cobertura de texto {cobertura:.1%}")
    else:
        texto = ""

    # 5. Regras
    if tipo == "NT":
        regras = regras_da_nt(dados)
    else:
        regras = dados.get("regras_validacao", [])
    info["regras"] = len(regras)

    if tipo == "NT" and texto:
        ids_pdf = set(PAT_ID_LINHA.findall(texto))
        ids_json = {r.get("id") for r in regras}
        so_pdf = sorted(ids_pdf - ids_json)
        info["ids_pdf"] = len(ids_pdf)
        if so_pdf:
            add(BLOQUEANTE, "REGRAS_AUSENTES",
                f"{len(so_pdf)} ID(s) do PDF fora do JSON: {', '.join(so_pdf[:8])}")

    # Nas NTs colchete desbalanceado conta como truncamento; nos MOCs so e aviso (exemplos do proprio PDF,
    # como '[Param: xxx)', sao desbalanceados de verdade)
    truncadas = [r for r in regras if mensagem_truncada(r.get("mensagem"), colchetes=(tipo == "NT"))]
    info["mensagens_truncadas"] = len(truncadas)
    if truncadas:
        ex = ", ".join(str(r.get("id")) for r in truncadas[:5])
        add(BLOQUEANTE, "MENSAGEM_TRUNCADA", f"{len(truncadas)}/{len(regras)} mensagens cortadas (ex.: {ex})")
    if tipo != "NT":
        colchetes = [r for r in regras if mensagem_truncada(r.get("mensagem"), colchetes=True)]
        if colchetes:
            ex = ", ".join(str(r.get("id")) for r in colchetes[:5])
            add(AVISO, "COLCHETES_DESBALANCEADOS", f"{len(colchetes)} mensagens com colchetes desbalanceados (ex.: {ex})")

    coladas = []
    for r in regras:
        for n in r.get("notas", []):
            if PAT_NOTA_COLADA.search((n.get("texto") or "").strip()):
                coladas.append(r.get("id"))
                break
    if coladas:
        add(AVISO, "NOTA_COM_FRAGMENTO", f"{len(coladas)} regra(s) com nota terminando em fragmento (ex.: {coladas[:5]})")

    # 6. Cronograma (NT)
    if tipo == "NT":
        cron = dados.get("cronograma", [])
        info["cronograma"] = len(cron)
        if not cron:
            add(BLOQUEANTE, "CRONOGRAMA_VAZIO", "cronograma vazio (ERRO_DATA_NAO_ENCONTRADA)")
        for motivo in cronograma_invalido(cron):
            add(BLOQUEANTE, "CRONOGRAMA_INVALIDO", motivo)

    # 7. Marcacoes
    stats = dados.get("estatisticas", {}).get("por_marcacao", {})
    if caminho_origem and caminho_origem.lower().endswith(".pdf"):
        fundos = fundos_coloridos(caminho_origem)
        total_json = stats.get("AMARELO", 0) + stats.get("VERDE", 0)
        if (fundos["AMARELO"] or fundos["VERDE"]) and total_json == 0:
            add(BLOQUEANTE, "MARCACAO_PERDIDA", "PDF tem fundo amarelo/verde mas o JSON nao tem nenhuma marcacao")
        elif not (fundos["AMARELO"] or fundos["VERDE"]) and total_json == 0:
            info["sem_destaque_no_pdf"] = True
        info["paginas_amarelo"] = sorted(fundos["AMARELO"])
        info["paginas_verde"] = sorted(fundos["VERDE"])

    return {"arquivo": nome, "tipo": tipo, "problemas": problemas, "info": info}


def listar_jsons():
    base = os.path.join(RAIZ, "catalogo")
    return sorted(glob.glob(os.path.join(base, "nt", "*", "*.json"))
                  + glob.glob(os.path.join(base, "moc", "*", "*.json")))


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    so_bloq = "--so-bloqueantes" in sys.argv
    arquivos = args or listar_jsons()
    if not arquivos:
        print("Nenhum JSON encontrado em catalogo/.")
        return 1

    total_bloq = 0
    total_aviso = 0
    for caminho in arquivos:
        r = verificar_json(caminho)
        bloq = [p for p in r["problemas"] if p[0] == BLOQUEANTE]
        avis = [p for p in r["problemas"] if p[0] == AVISO]
        total_bloq += len(bloq)
        total_aviso += len(avis)
        estado = "FALHA" if bloq else ("AVISO" if avis else "OK")
        resumo = ", ".join(f"{k}={v}" for k, v in r["info"].items() if not isinstance(v, list))
        print(f"[{estado:5}] {r['tipo']:3} {r['arquivo']}  ({resumo})")
        for nivel, codigo, texto in r["problemas"]:
            if so_bloq and nivel != BLOQUEANTE:
                continue
            print(f"         {nivel}: {codigo} - {texto}")

    print(f"\nArquivos: {len(arquivos)} | bloqueantes: {total_bloq} | avisos: {total_aviso}")
    return 1 if total_bloq else 0


if __name__ == "__main__":
    sys.exit(main())
